from pathlib import Path
from datetime import date
import argparse, json, shutil

from jinja2 import Environment, FileSystemLoader
from docx import Document
from docx.shared import Pt, RGBColor
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from utils import load_yaml, filter_items, has_any_tag, ensure_dir

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'data'
TARGETS = ROOT / 'targets'
TEMPLATES = ROOT / 'templates'
ASSETS = ROOT / 'assets'
OUTPUT = ROOT / 'output'

TARGET_IDS = ['full', 'dotnet', 'delphi', 'web', 'ai']
LANGS = ['en', 'uk']

FONT_CANDIDATES = [
    ('DejaVuSans', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
     'DejaVuSans-Bold', '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'),
    ('DejaVuSans', 'C:/Windows/Fonts/DejaVuSans.ttf',
     'DejaVuSans-Bold', 'C:/Windows/Fonts/DejaVuSans-Bold.ttf'),
    ('ArialUnicode', 'C:/Windows/Fonts/arial.ttf',
     'ArialUnicode-Bold', 'C:/Windows/Fonts/arialbd.ttf'),
]
_fonts = {'regular': 'Helvetica', 'bold': 'Helvetica-Bold'}


def register_fonts():
    """Register a Unicode font; fall back to a system one so local builds work too."""
    for reg_name, reg_path, bold_name, bold_path in FONT_CANDIDATES:
        if Path(reg_path).exists() and Path(bold_path).exists():
            pdfmetrics.registerFont(TTFont(reg_name, reg_path))
            pdfmetrics.registerFont(TTFont(bold_name, bold_path))
            _fonts['regular'], _fonts['bold'] = reg_name, bold_name
            return
    print('warning: no Unicode TTF found, PDF falls back to Helvetica (Latin only)')


# ---------------------------------------------------------------- helpers

def t(value, lang):
    """Read a {uk, en} value, tolerating plain strings."""
    if value is None:
        return ''
    if isinstance(value, str):
        return value
    return value.get(lang) or value.get('en') or value.get('uk') or ''


def norm(value):
    """Normalize a value to a {uk, en} dict."""
    if isinstance(value, dict):
        return {'uk': value.get('uk') or value.get('en') or '', 'en': value.get('en') or value.get('uk') or ''}
    return {'uk': value or '', 'en': value or ''}


def filter_groups(groups, tags, full_mode):
    """Keep skill groups and the items inside them that match the target."""
    out = []
    for g in groups:
        items = [dict(s, name=norm(s['name'])) for s in g.get('items', [])
                 if full_mode or has_any_tag(s.get('tags', []), tags)]
        if not items:
            continue
        if not (full_mode or has_any_tag(g.get('tags', []), tags)):
            continue
        out.append({'id': g['id'], 'label': norm(g['label']), 'items': items, 'primary': g.get('primary')})
    return out


def by_profile(items, target_id):
    """Projects whose `primary` names this profile come first; order inside each half is kept."""
    return sorted(items, key=lambda x: 0 if target_id in (x.get('primary') or []) else 1)


def bullets_for(exp, target_id):
    """Pick and order bullets per profile via `bullet_order`; fall back to `default`, then to all."""
    order = exp.get('bullet_order') or {}
    idx = order.get(target_id) or order.get('default')
    if not idx:
        return exp
    picked = {lang: [lst[i] for i in idx if i < len(lst)] for lang, lst in exp['bullets'].items()}
    return dict(exp, bullets=picked)


def is_detailed(exp, target_id, full_mode):
    """`detail_in` lists the profiles where a position gets bullets; elsewhere it folds into one line."""
    if 'detail_in' not in exp:
        return True
    return target_id in exp['detail_in']


def earlier_text(items, lang, present_word):
    """Older or short positions in one line: company — title (years)."""
    parts = []
    for e in items:
        end = present_word if e['end'] == 'present' else e['end'][:4]
        parts.append(f"{t(e['company'], lang)} — {t(e['title'], lang)} ({e['start'][:4]}–{end})")
    return ' · '.join(parts)


def period_text(item, lang, present_word):
    end = item['end']
    if end == 'present':
        end = present_word
    return f"{item['start']} — {end}"


# ---------------------------------------------------------------- documents

def doc_sections(bundle, lang):
    """Shared content model for md / docx / pdf."""
    p = bundle['profile']
    present = t(p['ui']['present'], lang)
    return {
        'name': t(p['name'], lang),
        'title': t(bundle['title_obj'], lang),
        'summary': t(bundle['summary_obj'], lang),
        'location': t(p['location'], lang),
        'contacts': [c['label'] for c in p['contacts']],
        'labels': dict({k: t(v, lang) for k, v in p['sections'].items()},
                       earlier_label=t(p['ui']['earlier'], lang) + ': '),
        'metrics': [f"{m['value']} {t(m['unit'], lang)} — {t(m['label'], lang)}".replace('  ', ' ')
                    for m in bundle['metrics']],
        'achievements': [t(a['text'], lang) for a in bundle['achievements']],
        'stack': [(t(g['label'], lang), [t(s['name'], lang) for s in g['items']])
                  for g in bundle['skill_groups']],
        'projects': [{
            'name': t(pr['name'], lang),
            'period': pr['period'].replace('present', present),
            'role': t(pr.get('role'), lang),
            'description': t(pr['description'], lang),
            'stack': pr.get('stack') or [],
            'link': (pr.get('link') or {}).get('href', ''),
        } for pr in bundle['projects']],
        'experiences': [{
            'company': t(e['company'], lang),
            'note': t(e.get('company_note'), lang),
            'title': t(e['title'], lang),
            'period': period_text(e, lang, present),
            'city': t(e['city'], lang),
            'bullets': e['bullets'][lang if lang in e['bullets'] else 'en'],
        } for e in bundle['experiences']],
        'earlier': earlier_text(bundle['earlier'], lang, present),
        'domain': [(t(d['term'], lang), t(d['note'], lang)) for d in bundle['domain']],
        'education': [{
            'specialty': t(e['specialty'], lang),
            'institution': t(e['institution'], lang),
            'period': e['period'],
            'kind': e.get('kind', 'degree'),
        } for e in bundle['education']],
        'languages': [(t(l['name'], lang), t(l['level'], lang)) for l in bundle['languages']],
        'side_projects': [(t(x['name'], lang), t(x['note'], lang), x['link']) for x in bundle['side_projects']],
    }


def build_docx(path, s):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10.5)

    doc.add_heading(s['name'], 0)
    p = doc.add_paragraph()
    p.add_run(s['title']).bold = True
    doc.add_paragraph(' · '.join(s['contacts'] + [s['location']]))
    doc.add_paragraph(s['summary'])

    if s['metrics']:
        doc.add_heading(s['labels'].get('summary_cards', 'At a glance'), level=1)
        for m in s['metrics']:
            doc.add_paragraph(m, style='List Bullet')

    doc.add_heading(s['labels']['strengths'], level=1)
    for a in s['achievements']:
        doc.add_paragraph(a, style='List Bullet')

    doc.add_heading(s['labels']['stack'], level=1)
    for label, items in s['stack']:
        par = doc.add_paragraph()
        par.add_run(f'{label}: ').bold = True
        par.add_run(', '.join(items))

    doc.add_heading(s['labels']['projects'], level=1)
    for pr in s['projects']:
        par = doc.add_paragraph()
        par.add_run(f"{pr['name']} ({pr['period']})").bold = True
        if pr['role']:
            doc.add_paragraph(f"{pr['role']}")
        doc.add_paragraph(pr['description'])
        if pr['stack']:
            doc.add_paragraph(', '.join(pr['stack']))
        if pr['link']:
            doc.add_paragraph(pr['link'])

    doc.add_heading(s['labels']['details'], level=1)
    for e in s['experiences']:
        par = doc.add_paragraph()
        head = f"{e['company']} — {e['title']}"
        par.add_run(head).bold = True
        doc.add_paragraph(f"{e['period']} · {e['city']}")
        for b in e['bullets']:
            doc.add_paragraph(b, style='List Bullet')
    if s['earlier']:
        doc.add_paragraph(f"{s['labels'].get('earlier_label', '')}{s['earlier']}")

    if s['domain']:
        doc.add_heading(s['labels']['domain'], level=1)
        for term, note in s['domain']:
            par = doc.add_paragraph(style='List Bullet')
            par.add_run(f'{term}: ').bold = True
            par.add_run(note)

    doc.add_heading(s['labels']['education'], level=1)
    for e in s['education']:
        doc.add_paragraph(f"{e['specialty']} — {e['institution']}" + (f" ({e['period']})" if e['period'] else ''), style='List Bullet')

    if s['side_projects']:
        doc.add_heading(s['labels']['side'], level=1)
        for name, note, link in s['side_projects']:
            doc.add_paragraph(f'{name} — {note} — {link}', style='List Bullet')

    doc.add_heading(s['labels']['languages'], level=1)
    for name, level in s['languages']:
        doc.add_paragraph(f'{name} — {level}', style='List Bullet')

    doc.save(path)


def build_pdf(path, s):
    register_fonts()
    reg, bold = _fonts['regular'], _fonts['bold']
    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    left, right = 46, 46
    max_width = width - left - right
    y = height - 52

    def space(n):
        nonlocal y
        y -= n

    def page_break_if_needed(need=40):
        nonlocal y
        if y < need:
            c.showPage()
            y = height - 52

    def wrap(text, font, size, avail):
        words, lines, cur = str(text).split(), [], ''
        for w in words:
            probe = (cur + ' ' + w).strip()
            if pdfmetrics.stringWidth(probe, font, size) <= avail:
                cur = probe
            else:
                if cur:
                    lines.append(cur)
                cur = w
        if cur:
            lines.append(cur)
        return lines or ['']

    def text_block(value, font=None, size=10, leading=13.5, indent=0, color=(0.1, 0.12, 0.13)):
        nonlocal y
        font = font or reg
        c.setFillColorRGB(*color)
        for line in wrap(value, font, size, max_width - indent):
            page_break_if_needed(52)
            c.setFont(font, size)
            c.drawString(left + indent, y, line)
            y -= leading

    def heading(value):
        nonlocal y
        page_break_if_needed(72)
        space(8)
        c.setFillColorRGB(0.1, 0.12, 0.13)
        c.setFont(bold, 11.5)
        c.drawString(left, y, value.upper())
        y -= 5
        c.setStrokeColorRGB(0.75, 0.77, 0.75)
        c.setLineWidth(0.6)
        c.line(left, y, width - right, y)
        y -= 13

    text_block(s['name'], bold, 22, 26)
    text_block(s['title'], reg, 12.5, 17, color=(0.28, 0.32, 0.31))
    text_block(' · '.join(s['contacts'] + [s['location']]), reg, 9.5, 13, color=(0.35, 0.39, 0.38))
    space(4)
    text_block(s['summary'], reg, 10, 14, color=(0.18, 0.2, 0.21))

    if s['metrics']:
        heading(s['labels'].get('summary_cards', 'At a glance'))
        for m in s['metrics']:
            text_block('• ' + m, reg, 10, 14)

    heading(s['labels']['strengths'])
    for a in s['achievements']:
        text_block('• ' + a, reg, 10, 14)

    heading(s['labels']['stack'])
    for label, items in s['stack']:
        text_block(f"{label}: {', '.join(items)}", reg, 9.5, 13)

    heading(s['labels']['projects'])
    for pr in s['projects']:
        text_block(f"{pr['name']} ({pr['period']})" + (f" — {pr['role']}" if pr['role'] else ''), bold, 10, 14)
        text_block(pr['description'], reg, 9.5, 13, indent=10, color=(0.24, 0.27, 0.27))
        if pr['stack']:
            text_block(', '.join(pr['stack']), reg, 9, 12, indent=10, color=(0.4, 0.44, 0.43))
        if pr['link']:
            text_block(pr['link'], reg, 9, 12, indent=10, color=(0.4, 0.44, 0.43))
        space(3)

    heading(s['labels']['details'])
    for e in s['experiences']:
        text_block(f"{e['company']} — {e['title']}", bold, 10, 14)
        text_block(f"{e['period']} · {e['city']}", reg, 9, 12.5, color=(0.4, 0.44, 0.43))
        for b in e['bullets']:
            text_block('• ' + b, reg, 9.5, 13, indent=10, color=(0.24, 0.27, 0.27))
        space(4)
    if s['earlier']:
        text_block(s['labels'].get('earlier_label', '') + s['earlier'], reg, 9, 12.5, color=(0.4, 0.44, 0.43))

    if s['domain']:
        heading(s['labels']['domain'])
        for term, note in s['domain']:
            text_block(f'• {term}: {note}', reg, 9.5, 13)

    heading(s['labels']['education'])
    for e in s['education']:
        text_block(f"• {e['specialty']} — {e['institution']}" + (f" ({e['period']})" if e['period'] else ''), reg, 9.5, 13)

    if s['side_projects']:
        heading(s['labels']['side'])
        for name, note, link in s['side_projects']:
            text_block(f'• {name} — {note}', reg, 9.5, 13)
            text_block(link, reg, 9, 12, indent=10, color=(0.4, 0.44, 0.43))

    heading(s['labels']['languages'])
    for name, level in s['languages']:
        text_block(f'• {name} — {level}', reg, 9.5, 13)

    c.save()


# ---------------------------------------------------------------- build

def build_target(target_id, env):
    master = load_yaml(DATA / 'master.yaml')
    profile = master['profile']
    experiences = load_yaml(DATA / 'experience.yaml')
    projects = load_yaml(DATA / 'projects.yaml')
    skill_groups = load_yaml(DATA / 'skills.yaml')['groups']
    education = load_yaml(DATA / 'education.yaml')['education']
    achievements = load_yaml(DATA / 'achievements.yaml')['achievements']
    domain = load_yaml(DATA / 'domain.yaml')['domain']
    languages = load_yaml(DATA / 'languages.yaml')['languages']
    side_projects = load_yaml(DATA / 'side_projects.yaml')['side_projects']
    target = load_yaml(TARGETS / f'{target_id}.yaml')

    full_mode = target_id == 'full'
    tags = target['include_tags']

    bundle = {
        'profile': profile,
        'title_obj': profile['titles'][target['title_key']],
        'summary_obj': profile['summary'][target['summary_key']],
        'metrics': filter_items(profile['metrics'], tags, full_mode)[:4],
        'achievements': by_profile(filter_items(achievements, tags, full_mode), target_id)[:6],
        'skill_groups': by_profile(filter_groups(skill_groups, tags, full_mode), target_id),
        'projects': by_profile(filter_items(projects, tags, full_mode), target_id)[:6],
        'experiences': by_profile([bullets_for(e, target_id) for e in filter_items(experiences, tags, full_mode)
                                   if is_detailed(e, target_id, full_mode)], target_id),
        # свёрнутая строка «Раніше» — хронология для любого профиля, без фильтра по тегам
        'earlier': [e for e in experiences if not is_detailed(e, target_id, full_mode)],
        'domain': filter_items(domain, tags, full_mode),
        # запись без tags показывается во всех профилях
        'education': [e for e in education if 'tags' not in e or full_mode or has_any_tag(e['tags'], tags)],
        'languages': languages,
        'side_projects': filter_items(side_projects, tags, full_mode),
    }

    out_dir = OUTPUT / target_id
    ensure_dir(out_dir)
    prefix = '../' if target_id != 'full' else ''
    page_title = f"{t(profile['name'], 'en')} — {t(bundle['title_obj'], 'en')}"
    js_data = {
        'profile': profile,
        'sections': profile['sections'],
        'ui': profile['ui'],
        'titleTextObj': bundle['title_obj'],
        'summaryTextObj': bundle['summary_obj'],
    }

    html = env.get_template('site.html.j2').render(
        target=target,
        page_title=page_title,
        title_text=t(bundle['title_obj'], 'en'),
        name_text=t(profile['name'], 'en'),
        summary_text=t(bundle['summary_obj'], 'en'),
        location=norm(profile['location']),
        contacts=profile['contacts'],
        metrics=[dict(m, unit=norm(m['unit']), label=norm(m['label'])) for m in bundle['metrics']],
        achievements=bundle['achievements'],
        skill_groups=bundle['skill_groups'],
        projects=bundle['projects'],
        experiences=bundle['experiences'],
        earlier={'uk': earlier_text(bundle['earlier'], 'uk', t(profile['ui']['present'], 'uk')),
                 'en': earlier_text(bundle['earlier'], 'en', t(profile['ui']['present'], 'en'))},
        domain=bundle['domain'],
        education=bundle['education'],
        languages=bundle['languages'],
        side_projects=[dict(x, name=norm(x['name']), note=norm(x['note'])) for x in bundle['side_projects']],
        root_prefix=prefix,
        assets_prefix=prefix,
        downloads_prefix='./',
        build_date=date.today().isoformat(),
        js_data=json.dumps(js_data, ensure_ascii=False),
    )
    (out_dir / 'index.html').write_text(html, encoding='utf-8')

    for lang in LANGS:
        s = doc_sections(bundle, lang)
        suffix = '' if lang == 'en' else f'.{lang}'
        md = env.get_template('resume.md.j2').render(s=s)
        (out_dir / f'resume{suffix}.md').write_text(md, encoding='utf-8')
        build_docx(out_dir / f'resume{suffix}.docx', s)
        build_pdf(out_dir / f'resume{suffix}.pdf', s)


def prepare_site():
    site = OUTPUT / 'site'
    if site.exists():
        shutil.rmtree(site)
    site.mkdir(parents=True, exist_ok=True)
    shutil.copytree(ASSETS, site / 'assets', dirs_exist_ok=True)
    shutil.copytree(OUTPUT / 'full', site, dirs_exist_ok=True)
    for tid in TARGET_IDS[1:]:
        shutil.copytree(OUTPUT / tid, site / tid, dirs_exist_ok=True)
    (site / '.nojekyll').write_text('', encoding='utf-8')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--target', choices=TARGET_IDS)
    ap.add_argument('--all', action='store_true')
    args = ap.parse_args()

    env = Environment(loader=FileSystemLoader(str(TEMPLATES)), trim_blocks=False, lstrip_blocks=False)

    if args.all:
        for tid in TARGET_IDS:
            build_target(tid, env)
        prepare_site()
    elif args.target:
        build_target(args.target, env)
    else:
        raise SystemExit('Use --target or --all')


if __name__ == '__main__':
    main()
